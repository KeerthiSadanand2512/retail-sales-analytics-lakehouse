"""Convert the UCI Online Retail workbook to UTF-8 CSV without filtering rows."""
import argparse
import csv
from datetime import datetime
import hashlib
import json
from pathlib import Path
from openpyxl import load_workbook

COLUMNS = ['InvoiceNo', 'StockCode', 'Description', 'Quantity', 'InvoiceDate', 'UnitPrice', 'CustomerID', 'Country']

def normalize(value):
    if value is None:
        return ''
    if isinstance(value, datetime):
        return value.strftime('%Y-%m-%d %H:%M:%S')
    return str(value)

def convert(source, target, limit=None):
    source, target = Path(source), Path(target)
    if limit is not None and limit <= 0:
        raise ValueError('limit must be positive')
    if source.resolve() == target.resolve() or target.exists():
        raise ValueError('Choose a new output path; existing files are never overwritten')
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = workbook['Online Retail']
        rows = sheet.iter_rows(values_only=True)
        if list(next(rows)) != COLUMNS:
            raise ValueError('Unexpected UCI workbook columns')
        target.parent.mkdir(parents=True, exist_ok=True)
        count = 0
        with target.open('x', encoding='utf-8', newline='') as handle:
            writer = csv.writer(handle, lineterminator='\n')
            writer.writerow(COLUMNS)
            for row in rows:
                if limit is not None and count >= limit:
                    break
                writer.writerow([normalize(v) for v in row])
                count += 1
    finally:
        workbook.close()
    return {'rows': count, 'sample': limit is not None,
            'source_sha256': hashlib.sha256(source.read_bytes()).hexdigest(),
            'csv_sha256': hashlib.sha256(target.read_bytes()).hexdigest(),
            'conversion': 'XLSX values to UTF-8 CSV; ISO timestamp text; blanks retained; no filtering/deduplication'}

if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source')
    parser.add_argument('target')
    parser.add_argument('--limit', type=int, help='First N rows for a learning sample, not a representative sample')
    args = parser.parse_args()
    print(json.dumps(convert(args.source, args.target, args.limit), indent=2))

